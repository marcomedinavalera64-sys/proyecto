from pypdf import PdfReader
from docx import Document
from pathlib import Path
p=Path('work')
r=PdfReader(r'C:\Users\marco\Downloads\Desafio_CIIN1205P.pdf')
s='\n'.join(f'\n--- PAGINA {i+1} ---\n'+x.extract_text() for i,x in enumerate(r.pages))
p.joinpath('desafio.txt').write_text(s,encoding='utf-8')
d=Document(r'C:\Users\marco\Downloads\Grupo1_Informe.docx')
t='\n'.join(x.text for x in d.paragraphs)
for tb in d.tables:
 t+='\nTABLA\n'+'\n'.join(' | '.join(c.text for c in row.cells) for row in tb.rows)
p.joinpath('informe.txt').write_text(t,encoding='utf-8')
print(s)
print('\n--- WORD ---\n'+t)
