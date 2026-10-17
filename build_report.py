from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import textwrap, json

ROOT=Path(__file__).resolve().parent
OUT=ROOT.parent/'outputs'
FONT='C:/Windows/Fonts/arial.ttf'
BOLD='C:/Windows/Fonts/arialbd.ttf'
MONO='C:/Windows/Fonts/consola.ttf'

def listing(name,text,width=104):
    lines=[]
    for line in text.splitlines():
        lines.extend(textwrap.wrap(line,width=width,replace_whitespace=False,drop_whitespace=False) or [''])
    im=Image.new('RGB',(1500,55+len(lines)*27),'#f5f6f7'); d=ImageDraw.Draw(im)
    font=ImageFont.truetype(MONO,23)
    for i,line in enumerate(lines):d.text((25,25+i*27),line,font=font,fill='#14202a')
    im.save(ROOT/name)

def flow(name,sale=False):
    im=Image.new('RGB',(1450,1650),'white'); d=ImageDraw.Draw(im)
    f=ImageFont.truetype(FONT,29); fb=ImageFont.truetype(BOLD,29)
    def box(y,txt,shape='process',x=720,w=650,h=110):
        a,b=x-w//2,y; c,e=x+w//2,y+h
        if shape=='end':d.rounded_rectangle((a,b,c,e),radius=50,outline='#243a44',width=4,fill='#edf2f3')
        elif shape=='decision':d.polygon([(x,b),(c,y+h//2),(x,e),(a,y+h//2)],outline='#243a44',fill='#edf2f3',width=4)
        elif shape=='io':d.polygon([(a+35,b),(c,b),(c-35,e),(a,e)],outline='#243a44',fill='white',width=4)
        else:d.rectangle((a,b,c,e),outline='#243a44',width=4)
        ls=txt.split('\n')
        for j,t in enumerate(ls):d.text((x,y+h/2+(j-(len(ls)-1)/2)*34),t,anchor='mm',font=f,fill='black')
    def arrow(points,label=None,loc=None):
        d.line(points,fill='#243a44',width=4)
        x,y=points[-1]; px,py=points[-2]
        if y>py:d.polygon([(x,y),(x-10,y-18),(x+10,y-18)],fill='#243a44')
        elif x>px:d.polygon([(x,y),(x-18,y-10),(x-18,y+10)],fill='#243a44')
        elif x<px:d.polygon([(x,y),(x+18,y-10),(x+18,y+10)],fill='#243a44')
        else:d.polygon([(x,y),(x-10,y+18),(x+10,y+18)],fill='#243a44')
        if label:d.text(loc,label,font=fb,fill='black')
    if not sale:
        nodes=[(25,'Inicio','end'),(200,'Cargar y validar archivo JSON','process'),(380,'¿Carga correcta?','decision'),(580,'Mostrar menú y leer opción','io'),(780,'¿Opción igual a 0?','decision'),(990,'1 Registrar     2 Buscar     3 Ordenar\n4 Inventario     5 Resumen diario','process'),(1210,'Mostrar resultado o error\nOpción desconocida → aviso','io')]
        for n in nodes:box(*n)
        for a,b in [(135,200),(310,380),(490,580),(690,780),(890,990),(1100,1210)]:arrow([(720,a),(720,b)])
        d.text((735,515),'Sí',font=fb,fill='black');d.text((735,925),'No',font=fb,fill='black')
        box(380,'Aviso y fin','end',x=1230,w=350)
        arrow([(1045,435),(1055,435)],'No',(1070,490))
        box(780,'Fin','end',x=1230,w=350)
        arrow([(1045,835),(1055,835)],'Sí',(1070,890))
        arrow([(720,1320),(720,1430),(210,1430),(210,635),(395,635)])
        d.text((240,1450),'El ciclo continúa hasta elegir salir.',font=f,fill='black')
    else:
        for n in [(25,'Registrar pedido','end'),(190,'Leer cliente, canal y productos','io'),(360,'Normalizar y validar entradas\nAgrupar cantidades repetidas','process'),(550,'¿Datos válidos y stock suficiente?','decision'),(760,'Copiar estado y calcular importes\nDescontar stock en la copia','process'),(960,'Asignar código y agregar pedido\nGuardar copia en archivo temporal','process'),(1160,'¿Se reemplazó el archivo?','decision'),(1390,'Actualizar memoria y mostrar constancia','io')]:box(*n)
        for a,b in [(135,190),(300,360),(470,550),(660,760),(870,960),(1070,1160),(1270,1390)]:arrow([(720,a),(720,b)])
        box(550,'Rechazar\nsin cambios','io',x=1230,w=350)
        arrow([(1045,605),(1055,605)],'No',(1090,680))
        box(1160,'Avisar fallo\nsin confirmar','io',x=1230,w=350)
        arrow([(1045,1215),(1055,1215)],'No',(1090,1290))
        d.text((735,690),'Sí',font=fb,fill='black');d.text((735,1310),'Sí',font=fb,fill='black')
        box(1540,'Volver al menú','end',h=85)
        arrow([(720,1500),(720,1540)])
        arrow([(1230,660),(1425,660),(1425,1582),(1045,1582)])
        arrow([(1230,1270),(1425,1270)])
    im.save(ROOT/name)

flow('flujo_general.png');flow('flujo_venta.png',True)
code=(ROOT/'prototipo/negocio.py').read_text(encoding='utf-8')
listing('codigo.png','\n'.join(code.splitlines()[55:86]),95)
listing('persistencia.png','\n'.join((ROOT/'prototipo/archivos.py').read_text(encoding='utf-8').splitlines()[-19:]),95)
listing('pruebas.png',(ROOT/'pruebas.txt').read_text(encoding='utf-8-sig'),100)
console=(ROOT/'consola.txt').read_text(encoding='utf-8')
listing('consola.png',console,98)

d=Document();sec=d.sections[0];sec.page_height=Cm(29.7);sec.page_width=Cm(21)
sec.top_margin=Cm(2);sec.bottom_margin=Cm(1.8);sec.left_margin=Cm(2.2);sec.right_margin=Cm(2.2)
for s in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    st=d.styles[s];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
    st.paragraph_format.space_after=Pt(7)
d.styles['Normal'].font.size=Pt(10.5);d.styles['Normal'].paragraph_format.line_spacing=1.12
d.styles['Title'].font.size=Pt(25);d.styles['Heading 1'].font.size=Pt(17);d.styles['Heading 2'].font.size=Pt(12)
footer=sec.footer.paragraphs[0];footer.alignment=2
r=footer.add_run();fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');r._r.addnext(fld)
def p(t,style=None):return d.add_paragraph(t,style)
def h(t):d.add_heading(t,2)
def page(t):d.add_page_break();d.add_heading(t,1)
def table(headers,rows,widths=None):
    tb=d.add_table(rows=1,cols=len(headers));tb.autofit=False
    for i,x in enumerate(headers):tb.rows[0].cells[i].text=x
    for row in rows:
        for i,x in enumerate(row):tb.add_row() if False else None
        cells=tb.add_row().cells
        for i,x in enumerate(row):cells[i].text=str(x)
    for ri,row in enumerate(tb.rows):
        pr=row._tr.get_or_add_trPr();ns=OxmlElement('w:cantSplit');pr.append(ns)
        if ri==0:rep=OxmlElement('w:tblHeader');pr.append(rep)
        for ci,c in enumerate(row.cells):
            if widths:c.width=Cm(widths[ci])
            tc=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'5');e.set(qn('w:color'),'D9D9D9');b.append(e)
            tc.append(b)
            sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8EDF0' if ri==0 else 'FFFFFF');tc.append(sh)
            mar=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');mar.append(e)
            tc.append(mar)
            for pp in c.paragraphs:
                pp.paragraph_format.space_after=Pt(3);pp.paragraph_format.line_spacing=1.02
                for rr in pp.runs:rr.font.size=Pt(9);rr.bold=ri==0
    p('')
    return tb
def pic(name,width=16):d.add_picture(str(ROOT/name),width=Cm(width))
def mono(t):
    q=p(t);q.paragraph_format.line_spacing=1.0
    for r in q.runs:r.font.name='Consolas';r.font.size=Pt(9)
def cap(t):q=p(t);q.paragraph_format.space_after=Pt(9);q.runs[0].italic=True;q.runs[0].font.size=Pt(9)

d.add_picture(str(ROOT/'source_images/image2.png'),width=Cm(4.2))
p('FACULTAD DE INGENIERÍA')
p('Carrera de Ingeniería de Sistemas Computacionales')
p('\n')
p('Sistema de pedidos e inventario para la Quesería Los Alpes','Title')
p('Informe técnico de proyecto final','Subtitle')
p('Fundamentos de Programación  CIIN1205P')
p('Grupo 1')
p('Docente: Pereda Cabanillas, Edwuard Jhonatan')
p('Cajamarca, Perú\n2026-2')
p('\nIntegrantes y porcentajes de participación\n________________________________________________\n________________________________________________\n________________________________________________\n________________________________________________\n________________________________________________')

page('1 Introducción y alcance del proyecto')
p('El proyecto propone un programa de consola para registrar los pedidos de la Quesería Los Alpes, comprobar la disponibilidad de sus productos y conservar las ventas en un archivo. La prioridad es que una venta aceptada tenga un importe correcto y que el inventario no se reduzca cuando el pedido se rechaza. A estas funciones se suman la búsqueda por código, el ordenamiento de pedidos y un resumen de la recaudación diaria.')
p('El caso comercial se desarrolla con autorización del docente, según lo indicado por el grupo. Se conserva la exigencia técnica del desafío CIIN1205P: analizar un proceso, comparar soluciones, diseñar algoritmos, implementar funciones y arreglos, trabajar con cadenas y archivos, y comprobar los resultados. Los expedientes del caso municipal se sustituyen por pedidos identificados mediante un código único. Este prototipo comercial no se presenta como una mesa de partes ni como una implementación de la normativa de gobierno digital.')
h('Objetivo general')
p('Desarrollar y verificar un prototipo en Python que permita registrar, validar, consultar, ordenar y almacenar pedidos, manteniendo la relación entre las ventas y las existencias de cada producto.')
h('Objetivos específicos')
for x in ['Describir el flujo de atención y distinguir las actividades que realiza el programa de las tareas físicas del personal.','Validar los datos del cliente, el canal de atención, los productos y las cantidades antes de modificar el inventario.','Calcular los importes con precisión decimal, asignar códigos únicos y conservar los datos entre sesiones.','Evaluar el comportamiento con casos normales, límites de stock y fallas de entrada o almacenamiento.']:p(x,'List Bullet')
h('Alcance y límites')
p('La versión incluye queso fresco, queso mantecoso, queso suizo y mantequilla. Permite varios productos por pedido y funciona en una computadora, con una sola instancia abierta. No incorpora reposición, cambios de precio, anulaciones, crédito, impuestos, facturación electrónica ni varios cajeros simultáneos. La constancia emitida sirve para control interno; no reemplaza un comprobante tributario.')
p('El informe base describe el negocio y aporta ejemplos de ventas. Esos ejemplos se utilizan como datos de demostración. Las fotografías, las notas de entrevista, las mediciones de campo y los datos de los integrantes se completarán en los anexos correspondientes.')

page('2 Proceso de atención y análisis del problema')
h('Descripción del caso')
p('El informe base sitúa a la Quesería Los Alpes en el jirón Junín 965, Cajamarca, y describe una atención presencial y telefónica. El personal recibe el pedido, consulta las existencias, anota la información en un cuaderno, prepara el producto, calcula el importe y registra la venta. La elección del caso se justifica porque estas operaciones permiten aplicar validaciones, cálculos, búsquedas y actualización de datos a un proceso comercial concreto.')
p('El mismo informe menciona consultas de inventario de tres a siete minutos y esperas de hasta 35 minutos en horas de mayor demanda. Se conservan como estimaciones del antecedente, no como mediciones verificadas ni como base para afirmar un porcentaje de mejora. La ficha de observación del anexo permite contrastarlas posteriormente.')
h('Secuencia del proceso')
table(['Paso','Actividad del personal','Apoyo del programa'],[['1','Recibir el pedido y conocer el canal.','Capturar cliente, canal, producto y cantidad.'],['2','Confirmar existencias físicas.','Consultar el saldo registrado y validar la cantidad.'],['3','Confirmar el pedido.','Calcular el total y guardar el registro.'],['4','Cortar, pesar y empaquetar.','No ejecuta tareas físicas ni verifica la balanza.'],['5','Cobrar y entregar.','Mostrar una constancia interna.'],['6','Revisar el cierre.','Consultar ventas del día e inventario restante.']],[1.1,7.2,8.3])
h('Causas del problema según el enfoque de las seis M')
table(['Dimensión','Causa o riesgo identificado'],[['Personal','La atención depende de que cada vendedor anote todos los datos.'],['Equipos','El cuaderno no realiza búsquedas automáticas ni cálculos.'],['Entorno','La concurrencia de clientes dificulta actualizar el registro a tiempo.'],['Datos','Los apuntes pueden quedar incompletos o resultar difíciles de leer.'],['Métodos','La consulta, el cálculo y la actualización se hacen por separado.'],['Gestión','No existe una conciliación automática entre pedidos y stock.']],[3,13.6])
p('El efecto común es la pérdida de confiabilidad del registro. El software puede prevenir errores de digitación y cálculo; la concordancia con el stock físico sigue requiriendo controles del personal.')

page('3 Requerimientos y modelo de datos')
table(['Entrada','Proceso','Salida'],[['Cliente y canal','Normalizar espacios y validar contenido.','Datos aceptados o mensaje de corrección.'],['Producto y cantidad en kg','Agrupar productos repetidos y comprobar stock.','Pedido válido o rechazo sin cambios.'],['Precio y stock del catálogo','Calcular subtotales y descontar cantidades.','Importe y saldo actualizado.'],['Código del pedido','Recorrer la lista y comparar códigos.','Detalle del pedido o aviso de no encontrado.'],['Lista de pedidos y fecha','Ordenar por código y filtrar ventas del día.','Listado ordenado y resumen diario.'],['Archivo JSON','Leer, validar y guardar registros.','Datos disponibles al reiniciar.']],[4.2,7.2,5.2])
h('Datos y reglas de validación')
table(['Dato','Tipo y regla'],[['Cliente','Cadena de 2 a 80 caracteres; letras, espacios, apóstrofo y guion. Se reducen espacios repetidos.'],['Canal','Cadena; presencial o telefónico, sin distinguir mayúsculas.'],['Producto','Identificador entero del 1 al 4. El nombre y precio se toman del catálogo.'],['Cantidad','Decimal positivo; hasta tres decimales y seis cifras enteras. Acepta coma o punto.'],['Stock y precio','Stock no negativo y precio positivo. Son datos del catálogo de demostración.'],['Código','Cadena PED-000001 a PED-999999; consecutivo generado por el sistema.'],['Fecha y total','Fecha local con zona horaria; total monetario con dos decimales.']],[3.2,13.4])
p('Cada pedido contiene una lista de líneas con producto, kilogramos, precio unitario y subtotal. Los importes se guardan como cadenas decimales en JSON para recuperar su valor sin convertirlos a números binarios de punto flotante. El estado reúne dos listas: productos y pedidos.')
h('Criterios de aceptación')
p('Un pedido válido se guarda y recibe un código; uno inválido conserva el estado anterior. La búsqueda devuelve el código solicitado, el ordenamiento no pierde registros y un reinicio recupera el último pedido confirmado. Si el archivo está dañado, la aplicación se detiene con un aviso y lo conserva para revisión.')

page('4 Alternativas y selección de la solución')
p('Se compararon dos soluciones algorítmicas implementables por el equipo y la opción de contratar un servicio externo. La comparación utiliza criterios técnicos; no supone que se haya solicitado una cotización.')
table(['Criterio','A Lista y búsqueda lineal','B Índice y búsqueda binaria','C Servicio externo'],[['Organización','Lista de pedidos y JSON.','Lista ordenada e índice por código.','Depende de la propuesta del proveedor.'],['Consulta','Recorre hasta n pedidos.','Hasta log₂ n pasos en lista ordenada.','Debe verificarse en una demostración.'],['Ordenamiento','Inserción explícita, fácil de explicar.','Mantener orden tras cada alta.','Depende del sistema contratado.'],['Recursos','Python y almacenamiento local.','Python y mayor cuidado al mantener índices.','Contratación, configuración y soporte.'],['Viabilidad académica','Alta para el nivel del curso.','Posible, con mayor complejidad.','Menor evidencia de programación propia.'],['Costo','Tiempo de desarrollo y mantenimiento.','Más tiempo de diseño y pruebas.','Cotización inicial y posibles pagos periódicos.']],[2.5,4.7,4.7,4.7])
h('Valoración de las alternativas propias')
p('Se asignó una escala de 1 a 5, donde 5 es más favorable. Los pesos y puntajes son una valoración de diseño para este prototipo, no resultados de una encuesta: cumplimiento 35 %, facilidad de explicación 30 %, recursos disponibles 20 % y mantenimiento 15 %.')
table(['Alternativa','Cumple','Explicable','Recursos','Mantener','Puntaje'],[['A','5','5','5','4','4,85'],['B','5','3','4','3','3,90']],[4.1,2.3,2.7,2.5,2.5,2.5])
p('Se seleccionó la alternativa A. Su estructura hace visible la lógica de búsqueda y ordenamiento y permite concentrar el trabajo en la integridad del pedido. El menor rendimiento del ordenamiento por inserción se acepta para un volumen académico reducido; no se afirma que sea adecuado para un historial comercial de gran tamaño.')
h('Contraste con tercerizar')
p('Una contratación podría incorporar soporte, múltiples usuarios y facturación, pero habría que evaluar el costo total, la propiedad de los datos y la exportación del historial. Sin una cotización no corresponde asignarle un importe. Para comparar económicamente se propone sumar instalación, capacitación, soporte y pagos recurrentes durante el mismo periodo. La necesidad inmediata del curso se cubre mejor con un prototipo propio que pueda explicarse y modificarse.')

page('5 Principios de programación y organización')
h('Relaciones matemáticas e invariantes')
p('Para cada producto, el subtotal se obtiene como cantidad × precio por kilogramo y se redondea a dos decimales. El total del pedido es la suma de sus subtotales. La regla de inventario es stock final = stock inicial − cantidad vendida, con la condición 0 < cantidad ≤ stock inicial. Si un producto aparece varias veces en el pedido, sus cantidades se suman antes de comprobar esa condición.')
p('Ejemplo: 15 kg de queso fresco a S/ 18 por kg generan S/ 270,00. Con 20 kg iniciales quedan 5 kg. Un pedido de 20 kg de mantecoso con solo 15 kg disponibles se rechaza y mantiene el saldo en 15 kg. En una compra de 5 kg de suizo y 5 kg de mantequilla, el total es 5 × 28 + 5 × 25 = S/ 265,00.')
h('Descomposición del sistema')
table(['Módulo','Responsabilidad y funciones'],[['main.py','Interacción de consola, menú y presentación de resultados.'],['negocio.py','numero, inicial, registrar, buscar, ordenar y resumen. Contiene las reglas de la venta.'],['archivos.py','cargar, validar y guardar. Mantiene el estado en JSON.'],['test_prototipo.py','Verifica resultados con un estado independiente y archivos temporales.']],[3.5,13.1])
p('Las funciones reciben parámetros explícitos. registrar devuelve el pedido confirmado; buscar devuelve un registro o None; ordenar devuelve una nueva lista. guardar escribe el archivo y no devuelve un resultado de negocio. Las variables de cálculo son locales. La ruta del archivo se define en la interfaz y se entrega a la función de persistencia.')
h('Estructuras del curso')
p('La secuencia se observa al capturar, validar, calcular y guardar. Los condicionales rechazan entradas y seleccionan opciones; los bucles mantienen el menú y recorren productos. Las listas representan arreglos dinámicos: se usan arreglos paralelos para construir el catálogo y listas de registros para los pedidos. Cada pedido contiene una lista de detalles, equivalente a una estructura tabular de varias filas. Las cadenas se limpian con strip, split y join; los códigos se normalizan con upper.')
h('Eficiencia y verificación')
p('Para n pedidos, la búsqueda lineal necesita como máximo n comparaciones. El ordenamiento por inserción tiene costo cuadrático O(n²) en el peor caso y lineal O(n) si la lista ya está ordenada. La copia y escritura del estado aumentan con el historial completo. La comprobación consiste en contrastar las salidas con resultados calculados de antemano y verificar que los rechazos no alteren los datos.')

page('6 Diagrama de flujo del programa')
pic('flujo_general.png',15.8)
cap('Figura 1. Flujo general del prototipo. El óvalo indica inicio o fin, el paralelogramo entrada o salida, el rectángulo proceso y el rombo decisión.')
p('La carga del archivo se realiza antes de mostrar el menú. Un archivo inexistente permite iniciar el catálogo de demostración; un archivo inválido impide continuar. Cada operación finaliza en el menú y los errores de entrada se informan sin cerrar la sesión.')

page('7 Diagrama de flujo del registro de pedidos')
pic('flujo_venta.png',15.8)
cap('Figura 2. Registro de un pedido. La actualización en memoria ocurre después del guardado correcto.')
p('El pedido se procesa sobre una copia del estado. Si una línea supera las existencias o falla el almacenamiento, la venta completa se rechaza. Esta secuencia evita descuentos parciales cuando un pedido contiene varios productos.')

page('8 Pseudocódigo del sistema')
mono('''INICIO
  INTENTAR cargar y validar el archivo
  SI no existe, crear catálogo inicial y lista vacía
  SI hay error de lectura o validación, informar y TERMINAR
  REPETIR
    Mostrar menú y leer opción
    SEGÚN opción
      1: Leer cliente, canal y líneas; REGISTRAR
      2: Leer código; BUSCAR y mostrar resultado
      3: ORDENAR copia de pedidos y mostrarla
      4: Mostrar catálogo y stock
      5: Filtrar fecha de hoy; contar pedidos y sumar totales
      0: Terminar
      OTRO: Informar opción inválida
    Si ocurre error de entrada o escritura, mostrar aviso
  HASTA opción = 0
FIN

REGISTRAR cliente, canal, líneas
  Normalizar nombre y canal; validar valores
  Rechazar lista de líneas vacía
  Agrupar las cantidades de cada producto
  Copiar el estado actual
  PARA cada producto solicitado
    Validar producto y cantidad positiva
    SI cantidad > stock, rechazar el pedido completo
    subtotal ← redondear(cantidad × precio, 2)
    stock de la copia ← stock − cantidad
    Agregar detalle y acumular subtotal
  FIN PARA
  código ← máximo código existente + 1
  Rechazar si se supera el límite de seis cifras
  Agregar pedido con fecha, cliente, canal, detalle y total
  GUARDAR la copia
  SI guardado exitoso, sustituir estado y devolver pedido

BUSCAR pedidos, código
  Normalizar código quitando espacios y usando mayúsculas
  PARA cada pedido
    SI su código coincide, devolver pedido
  FIN PARA
  Devolver no encontrado''')

page('9 Pseudocódigo de ordenamiento y persistencia')
mono('''ORDENAR pedidos
  copia ← lista de pedidos
  PARA i desde 1 hasta longitud(copia) − 1
    actual ← copia[i]; j ← i − 1
    MIENTRAS j ≥ 0 Y copia[j].código > actual.código
      copia[j + 1] ← copia[j]
      j ← j − 1
    FIN MIENTRAS
    copia[j + 1] ← actual
  FIN PARA
  Devolver copia

GUARDAR ruta, estado
  Validar estructura, códigos e importes
  Crear archivo temporal en la misma carpeta
  Escribir estado completo como JSON en UTF-8
  Vaciar búfer y sincronizar escritura
  Reemplazar archivo de destino por el temporal
  Eliminar temporal sobrante si ocurre un error

CARGAR ruta
  SI archivo no existe, devolver estado inicial
  Leer JSON y validar estructura y datos
  SI contenido inválido, informar y conservar archivo
  Devolver estado validado''')
h('Elección del formato de archivo')
p('Se eligió JSON de texto porque conserva la relación entre productos, pedidos y detalles y permite inspeccionar el archivo durante las pruebas [2]. Un CSV simplifica las filas, pero obliga a separar encabezados y detalles o a repetir datos. Un archivo binario puede reducir legibilidad y requiere herramientas específicas para su revisión. Para este prototipo, la facilidad de verificar el contenido pesa más que reducir el tamaño.')
p('El guardado utiliza un temporal en la misma carpeta y os.replace para sustituir el destino [3]. Esto reduce el riesgo de dejar un archivo a medio escribir. No equivale a un sistema de respaldo ni resuelve escrituras simultáneas: se requiere una sola instancia, permisos de escritura y copias de seguridad independientes.')
h('Condiciones que deben mantenerse')
p('Los códigos son únicos, ningún stock es negativo y cada total coincide con la suma de detalles. Un fallo de lectura no crea silenciosamente un inventario nuevo. Al guardar, la memoria solo se actualiza después de confirmar la escritura. El catálogo de demostración no representa las existencias físicas actuales de la quesería.')

page('10 Implementación del registro y control de errores')
pic('codigo.png',16.4)
cap('Figura 3. Reproducción del fragmento implementado en negocio.py. El archivo completo se entrega con el prototipo.')
p('La función primero valida el cliente y agrupa las cantidades; el fragmento muestra la verificación de existencias, el cálculo y la preparación de la venta. El uso de Decimal desde cadenas evita introducir aproximaciones binarias al representar importes [1]. Se aplica ROUND_HALF_UP al subtotal de cada producto y después se suman los subtotales.')
p('El stock se modifica en nuevo, que es una copia. guardar(nuevo) debe terminar correctamente antes de reemplazar estado. Si falla, la excepción se comunica a la interfaz y el inventario anterior permanece intacto. El código se calcula con el mayor identificador existente, de modo que continúa al volver a abrir el archivo.')
h('Decisiones de legibilidad')
p('Los nombres de las funciones expresan una acción, las reglas comerciales están separadas del menú y los comentarios explican las decisiones menos evidentes. Se evita mezclar la preparación física del producto con una función del programa. Tampoco se confunde una lista en memoria con persistencia: esta última exige escribir y volver a leer un archivo.')

page('11 Archivos y evidencia de ejecución')
pic('persistencia.png',16.4)
cap('Figura 4. Guardado implementado en archivos.py con validación y reemplazo mediante un temporal.')
h('Uso del prototipo')
p('Para iniciar, se ejecuta python main.py desde la carpeta del proyecto. La opción 1 solicita cliente, canal y líneas del pedido; el producto 0 termina la captura. Las opciones 2 y 3 buscan y ordenan; la 4 muestra existencias y la 5 calcula el resumen de la fecha local. Cada venta válida se guarda automáticamente en datos/queseria.json.')
p('El archivo de datos incluye información comercial, por lo que se excluye del repositorio mediante .gitignore. Los ejemplos de prueba usan nombres de demostración. No se solicita DNI, dirección ni teléfono. El programa no incorpora permisos por usuario ni cifrado y debe mantenerse en un entorno de prueba controlado.')
h('Control de versiones y trabajo del equipo')
p('Git permite registrar cambios locales y GitHub alojar el repositorio compartido. Para el trabajo grupal se propone separar las tareas de validación, persistencia, interfaz y pruebas; cada integrante debe crear su rama con su identidad real, revisar los cambios y registrar commits descriptivos. El historial se consulta con git log --oneline --graph --all [5].')
p('Las capturas deben mostrar el historial fechado, los cambios de cada aporte y el repositorio remoto. El enlace y las evidencias de participación quedan reservados en el anexo C. Un conjunto de archivos entregado no demuestra por sí mismo trabajo progresivo ni colaboración; ese requisito se acredita con el historial real.')

page('12 Casos de prueba y resultados')
p('Se ejecutaron 18 pruebas automatizadas en Python. Los ejemplos provienen del informe base y se complementan con datos sintéticos para comprobar límites y errores. Cada prueba inicia con un estado independiente; por ello, los saldos de un caso no se trasladan al siguiente. No se presentan estas ejecuciones como pruebas de campo con clientes reales.')
table(['Caso','Entrada o condición','Resultado comprobado'],[['01','Daniel; fresco 15 kg a S/ 18.','S/ 270,00 y stock final 5 kg.'],['02','María; mantecoso 20 kg; stock 15.','Rechazo; estado sin cambios.'],['03','Yandel; suizo 5 kg y mantequilla 5 kg.','Total S/ 265,00.'],['04','Nombre en blanco.','Rechazo por nombre inválido.'],['05','Cantidad de −5 kg.','Rechazo por formato inválido.'],['06','Luis; suizo 3 kg; stock cero.','Rechazo por falta de stock.'],['07','Fresco 20 kg; stock 20 kg.','Aceptado; stock final cero.'],['08','0, letras, NaN, infinito y 1,0001.','Todas las entradas rechazadas.'],['09','Código existente con espacios y uno ausente.','Encuentra el primero; el segundo no.'],['10','Tres códigos en orden inverso.','Orden ascendente sin alterar el origen.'],['11','Guardar venta y leer el archivo de nuevo.','Datos recuperados iguales al estado.'],['12','Archivo con JSON mal formado.','Rechazo; archivo original conservado.'],['13','Fallo de escritura simulado.','Sin venta confirmada ni cambio de stock.'],['14','Fresco 15 kg y 10 kg en el mismo pedido.','Agrupa 25 kg; rechaza por stock.'],['15','Fresco 0,125 kg.','Aceptado; total S/ 2,25.'],['16','Resumen de fecha de venta y fecha sin ventas.','S/ 18,00 en la primera; cero pedidos en la otra.'],['17','Carrito vacío, producto 99 y canal web.','Rechazo en los tres escenarios.'],['18','Archivo con total manipulado a 999.','Rechazo por inconsistencia del total.']],[1.1,7.3,8.2])
p('Resultado global: 18 pruebas aprobadas, sin fallas en esta ejecución. La verificación cubre las reglas indicadas; no acredita rendimiento con grandes volúmenes, seguridad para producción ni funcionamiento simultáneo de varias cajas.')

page('13 Evidencias de pruebas y funcionamiento')
pic('pruebas.png',16.3)
cap('Figura 5. Salida real capturada de python -m unittest -v. Los archivos temporales se eliminan al terminar cada prueba.')
h('Comprobación de una venta en consola')
mono('''PED-000001 | Daniel | Total S/ 270.00
Queso fresco: 15 kg x S/ 18 = S/ 270.00
Inventario después de la venta
1 Queso fresco: 5 kg | S/ 18/kg
2 Queso mantecoso: 15 kg | S/ 22/kg
3 Queso suizo: 12 kg | S/ 28/kg
4 Mantequilla: 25 kg | S/ 25/kg
Resumen del día: 1 pedido | S/ 270.00''')
cap('Extracto de la ejecución capturada. La captura completa de salida de consola se conserva en el paquete del prototipo.')
h('Interpretación de los resultados')
p('El caso normal confirma simultáneamente importe y saldo. El caso de producto repetido comprueba un riesgo que no aparece al validar cada línea por separado. La simulación de fallo de escritura muestra que no se confirma una venta que no pudo guardarse. La revisión del total al cargar evita aceptar un archivo con importes internamente inconsistentes.')

page('14 Aprendizaje y reflexión técnica y ética')
h('Estrategias de aprendizaje aplicables al proyecto')
p('El trabajo se organizó a partir de ejemplos pequeños: una venta, una consulta y una escritura de archivo. Esta separación permite comparar cada salida con un resultado esperado antes de integrar el menú. La lectura de documentación oficial se concentra en la necesidad concreta: precisión de importes, representación del estado y pruebas repetibles. Los casos de prueba se conservan para volver a ejecutarlos después de cada cambio.')
table(['Necesidad','Estrategia y aplicación','Evidencia'],[['Importes precisos','Consultar Decimal y crear valores desde texto [1].','Caso 15 y cálculo de subtotales.'],['Conservar ventas','Estudiar JSON y separar lectura y escritura [2].','Caso 11 y módulo archivos.py.'],['Evitar confirmaciones fallidas','Revisar reemplazo de archivos y trabajar sobre una copia [3].','Caso 13.'],['Comprobar cambios','Usar unittest con datos aislados [4].','18 pruebas ejecutables.']],[3.7,8.1,4.8])
p('Estos recursos permiten ampliar el uso básico de funciones, listas y archivos. El equipo deberá contrastarlos con los temas efectivamente vistos en clase antes de declarar cuáles constituyen conocimiento nuevo. La evidencia individual debe describir qué comprendió cada integrante, qué cambió y cómo lo verificó; no basta con adjuntar una captura.')
h('Reflexión técnica')
p('La mejora principal frente a la propuesta inicial es conservar el estado y tratar cada pedido como una operación completa. La simplicidad ayuda a revisar la lógica, pero también impone límites: el archivo crece con las ventas, la búsqueda recorre la lista y no hay coordinación entre varias cajas. La siguiente versión debería incorporar reposición y anulaciones trazables antes de utilizarse para controlar un negocio real.')
h('Reflexión ética y social')
p('Un saldo incorrecto puede hacer que el personal prometa un producto que no tiene, mientras que un importe mal calculado afecta al cliente o al negocio. Por eso se rechaza el pedido completo ante una inconsistencia y se informa el error. El ahorro de tiempo es un beneficio esperado, pero no se afirma como medido hasta contar con observaciones comparables.')
p('Los nombres y las compras no deben publicarse en el repositorio. Las fotografías de campo deben tomarse con autorización y ocultar datos de clientes. Asimismo, la elaboración asistida por IA debe declararse: el grupo conserva la responsabilidad de revisar el programa y explicar sus decisiones durante la sustentación.')

page('15 Conclusiones y recomendaciones')
h('Conclusiones')
p('El prototipo cubre el registro, la validación, la búsqueda, el ordenamiento y el almacenamiento de pedidos para el caso comercial autorizado. Los cálculos y las restricciones de stock se expresan mediante reglas comprobables. Las listas, funciones, cadenas, estructuras de control y archivos tienen una función concreta en la solución.')
p('Las 18 pruebas ejecutadas comprobaron escenarios normales, límites y errores, incluidos el guardado y la recuperación del estado. El resultado respalda el comportamiento de esta versión con los casos documentados. No permite concluir que la aplicación esté lista para varias cajas ni que haya reducido tiempos reales de atención.')
p('La información de campo y el historial de colaboración complementan la evidencia técnica. El prototipo, las pruebas y los diagramas permiten revisar la solución; las fotografías fechadas, entrevistas y aportes en GitHub deben acreditar el proceso real del grupo.')
h('Recomendaciones técnicas y conocimiento complementario')
p('Se recomienda conservar Decimal para los importes [1] y documentar la regla de redondeo. Su aplicación evita mezclar montos monetarios con cálculos aproximados innecesarios. La prueba con 0,125 kg muestra un caso decimal concreto; conviene añadir casos de redondeo a medio céntimo al modificar precios.')
p('El guardado mediante temporal y reemplazo [3] amplía la escritura directa de archivos: primero prepara el estado completo y después sustituye el destino. Debe complementarse con respaldos y una prueba de recuperación, porque un reemplazo correcto no protege frente a borrado accidental o falla del equipo.')
p('Las pruebas con unittest [4] deben ejecutarse después de cada cambio. Antes de incorporar reposiciones o anulaciones, se deben definir sus reglas y verificar que stock, historial y total permanezcan consistentes. Para varios usuarios será necesario revisar el modelo de almacenamiento y la concurrencia, no solo cambiar la apariencia de la interfaz.')
h('Verificación de campo y preparación de la sustentación')
p('Se recomienda medir varias atenciones con un mismo criterio de inicio y fin, registrar el número de observaciones y comparar periodos semejantes. En la exposición, el equipo debe demostrar una venta válida, un rechazo por stock, una búsqueda y la recuperación de datos tras reiniciar. También debe explicar por qué la venta se guarda antes de actualizar la memoria.')
p('Para la entrega final se adjuntan el informe editable, su versión PDF, una presentación de hasta 12 diapositivas y el código con pruebas. El grupo completará los campos reservados de autoría, campo y repositorio, y adjuntará la declaración oficial de IA disponible en el aula virtual.')

page('16 Referencias')
p('Documentos del curso y antecedente')
p('Universidad Privada del Norte. (s. f.). Evaluación final del curso Fundamentos de Programación CIIN1205P. Desafío Del archivador al algoritmo, un prototipo para la mesa de partes digital. Documento proporcionado para la actividad.')
p('Grupo 1. (2026). Informe de proyecto final sobre la Quesería Los Alpes. Documento base proporcionado para la revisión. Sus ejemplos de precios y existencias se utilizan como datos de demostración.')
p('Fuentes técnicas consultadas')
for t in ['[1] Python Software Foundation. (s. f.). decimal — Decimal fixed-point and floating-point arithmetic. https://docs.python.org/3/library/decimal.html','[2] Python Software Foundation. (s. f.). json — JSON encoder and decoder. https://docs.python.org/3/library/json.html','[3] Python Software Foundation. (s. f.). os — Miscellaneous operating system interfaces. Apartado os.replace. https://docs.python.org/3/library/os.html#os.replace','[4] Python Software Foundation. (s. f.). unittest — Unit testing framework. https://docs.python.org/3/library/unittest.html','[5] Git. (s. f.). git-log Documentation. https://git-scm.com/docs/git-log']:p(t)
p('Fecha de consulta: 14 de octubre de 2026. El prototipo se verificó con el intérprete de Python disponible en el entorno de trabajo y utiliza únicamente su biblioteca estándar.')
h('Correspondencia con el desafío autorizado')
p('Se aplican al proceso comercial las exigencias de análisis, alternativas, diseño, programación, pruebas y aprendizaje. La referencia municipal y sus obligaciones normativas pertenecen al caso original del desafío y no se trasladan a la quesería. La adaptación del caso responde a la autorización del docente comunicada por el grupo.')

page('Anexo A Instrumentos para la práctica de campo')
p('Ficha de observación y bitácora fotográfica. Completar durante la visita con información verificable.')
for t in ['Institución y dirección: ______________________________________________','Fecha y horario: ___________________________________________________','Responsable que autoriza la visita: ___________________________________','Integrantes presentes: ______________________________________________']:p(t)
table(['Registro','Qué documentar','Dato o evidencia'],[['F01','Recepción del pedido y formato utilizado.','Fecha y archivo: __________________'],['F02','Consulta de existencias, sin datos personales.','Fecha y archivo: __________________'],['F03','Registro y cálculo del importe.','Fecha y archivo: __________________'],['F04','Actualización y revisión al cierre.','Fecha y archivo: __________________']],[1.6,8.1,6.9])
p('Para cada fotografía: registrar fecha, lugar, actividad, autorización y descripción. Ocultar nombres, teléfonos, comprobantes y cualquier dato identificable de clientes.')
table(['Atención','Inicio','Fin','Duración','Incidencia'],[['1','________','________','________','________________'],['2','________','________','________','________________'],['3','________','________','________','________________']],[2,2.5,2.5,2.5,7.1])
p('Criterio de medición: inicio cuando el cliente comunica su pedido; fin cuando se confirma su aceptación o rechazo. Registrar aparte preparación y entrega. Duración = hora final − hora inicial. Ampliar la ficha si se observan más atenciones.')
h('Guía breve de entrevista')
for t in ['¿Qué datos se anotan y cuáles suelen faltar?','¿Cómo se confirma el stock y quién lo actualiza?','¿Qué ocurre cuando un pedido incluye varios productos?','¿Qué errores se repiten al cobrar o cerrar la jornada?','¿Cómo se consultan pedidos anteriores y cómo se conservan?','¿Qué cambio sería más útil y quién validaría el prototipo?']:p(t,'List Number')
p('Nombre o cargo y consentimiento: ____________________________________\nRespuestas y notas: __________________________________________________\n__________________________________________________________________')

page('Anexo B Declaración de uso de inteligencia artificial')
p('Texto de apoyo para trasladar al formato oficial del aula virtual. Esta hoja no reemplaza el formato institucional obligatorio ni una firma de los integrantes.')
h('Herramienta y finalidad')
p('Se utilizó ChatGPT en el entorno Codex para revisar los documentos proporcionados, identificar los requisitos del desafío, reorganizar la redacción y elaborar el código del prototipo, diagramas, pruebas y materiales de sustentación. El apoyo comprendió generación de contenido y programación, además de corrección de estilo.')
h('Verificaciones realizadas durante la preparación')
p('Se contrastó el contenido con el desafío y el informe base, se verificaron los cálculos de los ejemplos y se ejecutaron 18 pruebas automatizadas sobre el programa. Los datos de ventas utilizados son ejemplos de demostración; no se generaron fotografías de campo, entrevistas ni aportes ficticios de integrantes.')
h('Revisión y responsabilidad del grupo')
p('Antes de firmar la declaración, cada integrante debe revisar el contenido, ejecutar o presenciar la demostración y comprender las funciones que presenta. Las correcciones propias y las verificaciones adicionales se registrarán en los espacios siguientes.')
p('Integrantes que revisaron: ___________________________________________\nFecha de revisión: __________________________________________________\nCambios realizados por el equipo: ____________________________________\n__________________________________________________________________\nComprobaciones adicionales: _________________________________________\n__________________________________________________________________\nFirmas según formato oficial: ________________________________________')
h('Registro individual de aprendizaje')
table(['Integrante','Conocimiento nuevo','Aplicación y evidencia'],[['____________','____________________','________________________'],['____________','____________________','________________________'],['____________','____________________','________________________'],['____________','____________________','________________________'],['____________','____________________','________________________']],[4.1,5.8,6.7])

page('Anexo C Control de entrega y evidencias del grupo')
table(['Requisito','Ubicación o estado'],[['Caso autorizado y justificación','Secciones 1 y 2.'],['Entradas, procesos y salidas','Sección 3.'],['Dos alternativas y contraste con tercerizar','Sección 4.'],['Principios matemáticos y de ingeniería','Sección 5.'],['Diagrama de flujo y pseudocódigo','Secciones 6 a 9.'],['Funciones, arreglos, cadenas y archivos','Secciones 5 y 10 a 11; código adjunto.'],['Resultados de pruebas ejecutadas','Secciones 12 y 13; 18 pruebas.'],['Reflexión técnica y ética','Sección 14.'],['Aprendizaje y conocimiento complementario','Secciones 14 y 15; registro individual en anexo B.'],['Observación, fotos fechadas y entrevista','Completar con evidencias reales en anexo A.'],['Pruebas con datos recogidos en campo','Añadir al validar la información de la visita.'],['Git, GitHub e historial progresivo','Completar enlace y capturas reales debajo.'],['Informe PDF y presentación','Archivos adjuntos con el nombre del grupo.'],['Declaración oficial de IA','Trasladar anexo B al formato del aula y adjuntarlo.'],['Integrantes y participación','Completar carátula con datos reales.']],[7.1,9.5])
p('Enlace al repositorio: ________________________________________________\nRama o versión entregada: ___________________________________________\nIdentificador del commit final: ________________________________________')
p('Adjuntar capturas legibles del historial fechado, aportes por integrante y vista del repositorio remoto. No publicar archivos con datos reales de clientes.')
p('Entrega: aula virtual, según el sílabo, semanas 15 y 16. Límite del desafío: 20 páginas de contenido sin contar carátula ni anexos y 12 diapositivas. Revisar que los archivos correspondan a la misma versión antes de subirlos.')

d.core_properties.title='Sistema de pedidos e inventario para la Quesería Los Alpes'
d.core_properties.author='Grupo 1'
d.core_properties.subject='Fundamentos de Programación CIIN1205P'
for el in d.styles.element.iter():
    if el.tag == qn('w:pBdr'):
        el.getparent().remove(el)
for el in d.element.iter():
    if el.tag == qn('w:pBdr'):
        el.getparent().remove(el)
d.save(OUT/'Grupo1_FundamentosDeProgramacion_EF.docx')
print('Documento creado')
